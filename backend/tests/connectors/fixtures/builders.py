"""Synthetic workbooks in the real shape of each state's list, written with openpyxl at test time.

Each builder reproduces the layout observed on 2026-10-27 (title rows, merged cells, repeated
column names, legends and footnotes below the data, Excel dates next to dates typed as text,
"N/A", several NPIs in one cell, placeholder end dates) with made-up people and identifiers.
No real row is ever copied here.
"""

import hashlib
import json
from collections.abc import Sequence
from datetime import date
from io import BytesIO

from openpyxl import Workbook

from tests.synthetic import make_npi

type Cell = str | int | float | date | None
type Rows = Sequence[Sequence[Cell]]

NPI_A = make_npi("100000001")
NPI_B = make_npi("100000002")
NPI_C = make_npi("100000003")
NPI_D = make_npi("100000004")
NPI_E = make_npi("100000005")
NPI_F = make_npi("100000006")
INVALID_TEN_DIGITS = "5203775500"
"""Ten digits that fail the NPI check digit, as some state cells hold."""


def workbook(
    rows: Rows,
    *,
    title: str = "Sheet1",
    merge: Sequence[str] = (),
    first_sheet: tuple[str, Rows] | None = None,
) -> bytes:
    """An .xlsx with ``rows`` on one sheet; ``first_sheet`` puts another sheet before it."""
    book = Workbook()
    sheet = book.active
    if first_sheet is not None:
        sheet.title = first_sheet[0]
        for row in first_sheet[1]:
            sheet.append(list(row))
        sheet = book.create_sheet()
    sheet.title = title
    for row in rows:
        sheet.append(list(row))
    for cells in merge:
        sheet.merge_cells(cells)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def nc_medicaid() -> bytes:
    """North Carolina: six preamble rows, header on row 7, owners in the last column."""
    return workbook(
        [
            ("STATE EXCLUDED PROVIDERS LIST - PUBLISHED SEPTEMBER 2026",),
            ("NC Medicaid - Office of Compliance and Program Integrity",),
            (),
            ('Exclusion applies only to the provider in "EXCLUDED ENTITY".\nOwners are not.',),
            (),
            ("Terminations through 08.31.2026",),
            (
                "EXCLUDED ENTITY",
                "NPI/ATYPICAL ID (EXCLUDED)",
                "EFFECTIVE DATE",
                "REASON FOR EXCLUSION  ",
                "CITY",
                "STATE",
                "ZIP CODE",
                "OWNERSHIP",
            ),
            (
                "SAMPLEVILLE CARE SERVICES LLC",
                int(NPI_A),
                date(2026, 2, 19),
                "FEDERAL",
                "SAMPLEVILLE\xa0",
                "NC",
                28501,
                "AVERY\xa0J TESTOWNER",
            ),
            (
                "JORDAN LEE EXAMPLEPERSON",
                int(NPI_B),
                date(2018, 10, 19),
                "OIG/STATE",
                "TESTTOWN",
                "NC",
                28512,
                "JORDAN LEE EXAMPLEPERSON",
            ),
            (
                "MOCK BILLING AGENCY",
                50345490,
                date(2020, 10, 20),
                "STATE",
                "FAKECITY",
                "GA",
                30361,
                "RILEY OWNER",
            ),
            (
                "CASEY (MAIDENMOCK) SAMPLEWOMAN",
                "N/A",
                "03/15/2021",
                "STATE",
                "TESTTOWN",
                "NC",
                28512,
                None,
            ),
            (
                "QUINN DE LA MOCK",
                int(NPI_C),
                date(2019, 1, 7),
                "STATE",
                "FAKECITY",
                "NC",
                27609,
                None,
            ),
            (
                "SAMPLE EAR NOSE AND THROAT",
                int(NPI_D),
                date(2017, 5, 2),
                "OIG",
                "TESTTOWN",
                "NC",
                28580,
                "OWNER ONE",
            ),
            (
                "SAMPLEVILLE CARE SERVICES LLC",
                int(NPI_A),
                date(2026, 2, 19),
                "FEDERAL",
                "SAMPLEVILLE",
                "NC",
                28501,
                "SECOND OWNER",
            ),
            ("RILEY MOCKPERSON", "N/A", "2019X345", "STATE", "TESTTOWN", "NC", 28512, None),
            (
                None,
                int(NPI_E),
                date(2020, 1, 1),
                "STATE",
                "TESTTOWN",
                "NC",
                28512,
                "ORPHAN OWNER",
            ),
        ],
        title="SEPTEMBER 2026",
    )


def md_medicaid() -> bytes:
    """Maryland: header on row 1, legend and list date below a blank row, empty rows after."""
    width = 26
    pad: tuple[None, ...] = (None,) * (width - 9)
    return workbook(
        [
            (
                "LAST NAME/ORGANIZATION",
                "FIRST NAME",
                "TYPE OF ENTITY/PROFESSION",
                "SANCTION TYPE",
                "NPI",
                "LICENSE NO",
                "TERMINATION/SANCTION DATE",
                "ADDRESS",
                "CITY/STATE/ZIP",
                *pad,
            ),
            (
                "Sampleville Medical Supply Company",
                None,
                "Business Entity",
                "MDH",
                None,
                None,
                date(2015, 10, 29),
                "146 Sample Drive Ste 100",
                "Testtown MD 21136",
            ),
            (
                "Testperson",
                "Avery",
                "MD",
                "LB",
                int(NPI_A),
                "D50296",
                date(2019, 12, 31),
                "7004 Mock Blvd, 101",
                "Fakecity, MD 21209",
            ),
            (
                "Exampledoc",
                "Jordan K",
                "RN",
                "HHS",
                f"{NPI_B}, {NPI_C}",
                None,
                date(2020, 8, 10),
                "19 Example Lane",
                "Sampleville, MD 21237",
            ),
            (
                "Mockperson",
                "Riley",
                "PT",
                "LB/HHS",
                f"{NPI_D};",
                "PT1234",
                date(2021, 4, 1),
                None,
                None,
            ),
            (
                "Samplewoman",
                "Casey",
                "Nurse",
                "MFVVU-A",
                None,
                None,
                date(2022, 2, 2),
                "699 Test Drive",
                "Glen Sample 21060",
            ),
            ("Fakename", "Quinn", "CNA", None, INVALID_TEN_DIGITS, None, None, None, "cics"),
            (),
            ("A = Abuse of adult/vulnerable individual",),
            ("LB = Licensing Board discipline (Suspended, surrendered, revoked license)",),
            ("HHS = Sanctioned by U.S. Department of Health and Human Services",),
            (),
            ("10.08.2026",),
            *((None,) * width for _ in range(5)),
        ],
        title="State_of_MD_Sanctioned_Provider",
    )


def vt_medicaid() -> bytes:
    """Vermont: a title row, header on row 2 with "State" twice, footnotes after a blank row."""
    open_ended = date(2382, 12, 31)
    return workbook(
        [
            ("#VALUE!", "Vermont Monthly List of For Cause Exclusions - October 2026"),
            (
                "VT Medicaid Provider ID ",
                "Provider NPI",
                "Provider Full Name",
                "Provider Type ",
                "Provider City",
                "State",
                "State",
                'Reason or "Cause" for Exclusion',
                "Provider Exclusion Status Effective Date",
                "Provider Exclusion Status End Date",
                "Reinstatement Date with Vermont Medicaid if Applicable",
            ),
            (
                "0001926",
                NPI_A,
                "TESTPERSON,AVERY J DMD                    ",
                "DENTIST          ",
                "SAMPLEVILLE       ",
                "VT",
                "22",
                "TERMINATED DUE TO SANCTIONS        ",
                date(2020, 2, 12),
                open_ended,
                None,
            ),
            (
                "0002001",
                None,
                "SAMPLE LAB WERKS INC",
                "LABORATORY",
                "FAKECITY",
                "NY",
                "07",
                "TERMED DUE TO MEDICARE/OTHER STATE EXCLUSION",
                date(2021, 6, 1),
                open_ended,
                None,
            ),
            (
                "0002002",
                NPI_B,
                "EXAMPLE,JORDAN LCMHC",
                "COUNSELOR",
                "TESTTOWN",
                "VT",
                "19",
                "TERMINATED DUE TO FAILURE OF CONTRACT TERMS",
                date(2019, 1, 1),
                date(2022, 1, 1),
                date(2023, 5, 1),
            ),
            (
                "0002003",
                f"{NPI_C}\n{NPI_D}",
                "MOCKCLINIC",
                "CLINIC",
                "TESTTOWN",
                "VT",
                "08",
                "FRAUD CONVICTION",
                "06/30/2024",
                open_ended,
                None,
            ),
            (),
            ('Names listed are of Providers Excluded for “Cause" from participation.',),
            ("Key",),
            ("Provider Exclusion Status Effective Date - The date the Exclusion began.",),
        ],
        title="Excluded Providers",
    )


def nd_medicaid() -> bytes:
    """North Dakota: a merged title row with the list date, then the header on row 2."""
    return workbook(
        [
            ("North Dakota Provider Exclusion List\n09/24/2026",),
            (
                "Provider Name",
                "Street Address",
                "City",
                "State",
                "Zip",
                "NPI",
                "License Number",
                "Provider Type ",
                "Exclusion Date",
                "Reason for Exclusion",
            ),
            (
                "Testperson, Avery",
                "2006 Sample St",
                "Sampleville",
                "ND",
                "58201",
                "N/A",
                "N/A",
                "Qualified Service Provider",
                date(2024, 8, 20),
                "Reciprocal to OIG Exclusion",
            ),
            (
                "Exampleperson (aka Mockname), Jordan",
                "1 Test Ave",
                "Fakecity",
                "ND",
                58501,
                int(NPI_A),
                "R12345",
                "Licensed Practical Nurse ",
                date(2022, 3, 1),
                "License revoked",
            ),
            (
                "Smiling Sample Care",
                "9 Example Rd",
                "Testtown",
                "MN",
                "56721-1818",
                "N/A",
                "N/A",
                "Qualified Service Provider",
                date(2023, 1, 5),
                "Billing fraud",
            ),
            (
                "Sampleman, Riley Jr",
                "3 Mock Ln",
                "Sampleville",
                "ND",
                585011234,
                "178037189",
                "N/A",
                "N/A",
                date(2021, 7, 9),
                "Reciprocal to OIG Exclusion",
            ),
        ],
        title="ND Provider Exclusion List",
        merge=("A1:J1",),
    )


def wa_hca() -> bytes:
    """Washington: title and update rows, a blank row, then the header on row 4."""
    return workbook(
        [
            ("Provider Termination and Exclusion List",),
            ("Updated 10/10/2026",),
            (),
            ("Name", "License #", "NPI # or P1 #", "Date of Exclusion", "Action"),
            (
                "Testperson, Avery M ARNP",
                "AP60000001",
                NPI_A,
                date(2023, 9, 21),
                "License Suspended",
            ),
            (
                "Sample Dentures NW\nMock Dental and Denture\nExampleowner, Jordan L. DN",
                "DN60406207",
                f"{NPI_B}\n{NPI_C}",
                date(2023, 9, 21),
                "Terminate Medicaid Authority",
            ),
            (
                "Mockdoc, Riley MD\ndba: Sample Pathology Labs",
                "MD00012345",
                "P1 - 1027986",
                date(2020, 2, 3),
                "OIG Exclusion",
            ),
            (
                "Samplewoman, Casey L. MD/\nSample Family Pediatrics",
                "MD00067890",
                f"{NPI_E}\n{NPI_F}",
                date(2022, 6, 1),
                "Terminated by CMS",
            ),
            ("Fakename, Casey", None, INVALID_TEN_DIGITS, "10/01/2020", "Revocation of GA Permit"),
            (
                "Example, Quinn PA-C",
                "PA001, PA002",
                NPI_D,
                date(2018, 4, 4),
                "Manditory Exclusion from Medicaid",
            ),
        ],
        title="Termination and exclusion list",
    )


def ia_medicaid() -> bytes:
    """Iowa: header on row 1 of a sheet named for the month, trailing empty columns."""
    pad = (None, None, None, None)
    return workbook(
        [
            (
                "Enrollment Type",
                "Specialty",
                "NPI",
                "Affiliated NPI",
                "Last Name",
                "First Name",
                "Legal Business Name",
                "Type of Sanction",
                "Effective Date",
                "Sanction End Date",
                "Eligible to Reapply Date",
                "Authority",
                "State License Type",
                "State License Number",
                *pad,
            ),
            (
                "Individual",
                "LPN",
                None,
                None,
                "Testperson",
                "Avery Beth",
                None,
                "Participation Suspension",
                date(2021, 3, 18),
                "Indefinite",
                None,
                "441 IAC 79.2(2)(a)",
                "LPN",
                "P57910",
                *pad,
            ),
            (
                "Organization",
                "Pharmacy",
                int(NPI_A),
                int(NPI_D),
                None,
                None,
                "Sample Pharmacy LLC",
                "Termination",
                date(2020, 5, 1),
                date(2024, 1, 1),
                date(2025, 1, 1),
                "441 IAC 79.2(2)(b)",
                None,
                None,
            ),
            (
                "Indvidual",
                "Physician",
                "X000000111",
                int(NPI_B),
                "Exampledoc",
                "Jordan",
                "Exampledoc Clinic",
                "Federal Exclusion",
                date(2019, 9, 9),
                "Federal Authority",
                None,
                "42 CFR 1001",
                "Physician",
                "MD-0001",
            ),
            (
                "individual",
                "Counselor",
                171092119,
                None,
                "Mockperson",
                "Riley",
                None,
                "Participation Suspension for Overpayment",
                "12/17/2025 - on payment plan",
                "2 Years",
                None,
                "441 IAC 79.2(2)(c)",
                None,
                None,
            ),
            (
                "Indivdual",
                "Nurse",
                int(NPI_C),
                "X000000222",
                "Samplewoman",
                "Casey",
                None,
                "Suspension",
                date(2026, 1, 2),
                date(2031, 6, 23),
                None,
                "441 IAC 79.2(2)(d)",
                "RN",
                "R-77",
            ),
        ],
        title="Sept 2026",
    )


def in_fssa() -> bytes:
    """Indiana: a title row, then the header on row 2 with a trailing space."""
    return workbook(
        [
            ("Provider Terminations (As of July 2026)",),
            (
                "Provider Name",
                "National Provider Identification (NPI)",
                "Service Location Address",
                "Termination Effective Date ",
            ),
            ("Avery Testperson", "N/A", "Indianapolis, IN", date(2015, 9, 17)),
            (
                "Jordan L. Exampledoc, MD",
                f"{NPI_A}\n{NPI_B}",
                "Sampleville, IN\nTesttown, IN\nFakecity, IN",
                date(2018, 3, 3),
            ),
            ("Sample Medical Laboratories", NPI_C, "Irvine, CA", date(2022, 8, 8)),
            ("Samplecorp", NPI_D, "Naperville, IL", date(2021, 1, 1)),
            (
                "Riley Mockperson, Podiatrist",
                NPI_E,
                "Gary, IN          Hammond, IN",
                date(2020, 2, 2),
            ),
            ("Casey Mc Example, DDS", NPI_F, "Phoenix AZ", "07/04/2019"),
        ],
    )


def mo_mmac() -> bytes:
    """Missouri: a title row, the header on row 2, one row per provider type (merged), a
    text date among Excel dates, a licence with another state's prefix."""
    return workbook(
        [
            ("PROVIDER SANCTION LIST - CURRENT THROUGH SEPTEMBER 2026",),
            (
                "PROVIDER NAME",
                "NPI",
                "TERM DATE",
                "LETTER DATE",
                "PROV TYPE/SPECIALTY",
                "LICENSE #",
                "TERMINATION REASON",
            ),
            (
                "SAMPLE CARE SERVICES LLC",
                int(NPI_A),
                date(2024, 6, 12),
                date(2024, 5, 13),
                "26-CDS",
                "N/A",
                "FAILURE TO USE EVV",
            ),
            (
                "SAMPLE CARE SERVICES LLC",
                int(NPI_A),
                date(2024, 6, 12),
                date(2024, 5, 13),
                "27-REASSESS ",
                "N/A",
                "FAILURE TO USE EVV",
            ),
            (
                "TESTPERSON, AVERY J JR",
                NPI_B,
                "02/28/19",
                date(2019, 1, 29),
                "20-PHYSICIAN",
                "KS 0412345, MO R1234",
                "OTHER STATE TERMINATION",
            ),
            (
                "A MOCK CDS",
                "N/A",
                date(2018, 9, 16),
                date(2018, 8, 16),
                "26-CDS ONLY",
                "N/A",
                "BAD DEBT",
            ),
            (
                "JORDAN K EXAMPLEDOC MD",
                int(NPI_C),
                date(2022, 11, 27),
                date(2022, 10, 27),
                "20-PHYSICIAN",
                2007012345,
                "OIG EXCLUSION",
            ),
            (
                "MOCKVILLE ADULT DAY CENTER",
                "N/A",
                date(2014, 7, 11),
                date(2014, 7, 9),
                "29- ADULT DAY HEALTH CARE",
                532,
                "FRAUD",
            ),
            (
                "MOCKVILLE ADULT DAY CENTER",
                "N/A",
                date(2014, 7, 11),
                date(2014, 7, 9),
                "58-HOME HEALTH AGENCY",
                756,
                "FRAUD",
            ),
        ],
        title="CURRENT",
    )


def ga_dch_oig() -> bytes:
    """Georgia: a title cell and a blank row above the header, YYYYMMDD dates as text and as
    integers, an Excel date, a made-up DOB, empty sheets after the data sheet."""
    rows: Rows = (
        (
            None,
            "Department of Community Health \nOffice of Inspector General\n"
            "List of Excluded Individuals and Entities as of 10/05/2026"
            "*Highlighted individuals recently added*",
        ),
        (),
        (
            "LAST NAME",
            "FIRST NAME",
            "MIDDLE NAME",
            "BUSINESS NAME",
            "GENERAL",
            "STATE",
            "SANCDATE",
            "NPI",
            "DOB",
        ),
        ("TESTPERSON", "AVERY", None, "", "MEDICAL PRACTICE, MD", "GA", "19820415", None),
        (
            "EXAMPLEDOC",
            "JORDAN ",
            "K",
            "",
            "NURSING PROFESSION",
            "AL, GA",
            20250630,
            int(NPI_A),
            date(1970, 1, 1),
        ),
        (None, None, None, "SAMPLE MEDICAL EQUIPMENT, INC", "DME COMPANY", "GA", 20111012),
        ("MOCKPERSON (MAIDENMOCK)", "RILEY", None, "", "REGISTERED NURSE", "GA", "20200101"),
        (
            "SAMPLEOWNER",
            "CASEY",
            None,
            "SAMPLE TRANSPORT CO",
            "BUS OWNER/EXEC",
            "GA",
            date(2026, 8, 27),
        ),
        ("FAKENAME", "QUINN", None, "", "PRIVATE CITIZEN", "GA", "20223108"),
    )
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    for row in rows:
        sheet.append(list(row))
    for name in ("Sheet4", "Sheet2", "Sheet3"):
        book.create_sheet(name)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def ca_dhcs_resource_show(body: bytes) -> bytes:
    """California: a CKAN ``resource_show`` answer (the real one's shape) describing ``body``."""
    return json.dumps(
        {
            "help": "https://data.chhs.ca.gov/api/3/action/help_show?name=resource_show",
            "success": True,
            "result": {
                "id": "48630a37-b5ba-4d3d-af54-e82b30e658a0",
                "url": (
                    "https://data.chhs.ca.gov/dataset/8f70e05d-baa5-4296-a5ab-36318d530670/"
                    "resource/48630a37-b5ba-4d3d-af54-e82b30e658a0/download/"
                    "suspended-ineligible-list-september-2026.csv"
                ),
                "last_modified": "2026-10-12T17:41:37.794529",
                "hash": hashlib.md5(body, usedforsecurity=False).hexdigest(),
                "size": len(body),
                "format": "CSV",
                "url_type": "upload",
            },
        }
    ).encode()
